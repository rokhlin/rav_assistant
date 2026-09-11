import io
import logging
from typing import Tuple, Optional, List
from pypdf import PdfReader
from docx import Document
from PIL import Image

logger = logging.getLogger(__name__)

class DocParser:
    @staticmethod
    def extract_from_pdf(file_bytes: bytes) -> Tuple[str, int]:
        """
        Extract text from PDF document.
        Returns (text, page_count).
        """
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            num_pages = len(reader.pages)
            text_parts = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_parts.append(f"--- Page {i+1} ---\n{page_text}")
            
            full_text = "\n\n".join(text_parts).strip()
            return full_text, num_pages
        except Exception as e:
            logger.error(f"Error parsing PDF: {e}")
            raise RuntimeError(f"Failed to read PDF: {e}")

    @staticmethod
    def extract_images_from_pdf(file_bytes: bytes) -> List[bytes]:
        """
        Extract embedded images from PDF pages (for scanned documents without text layer).
        """
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            images = []
            for page in reader.pages:
                try:
                    for img in page.images:
                        images.append(img.data)
                except Exception as img_err:
                    logger.debug(f"Error extracting image from page: {img_err}")
            return images
        except Exception as e:
            logger.warning(f"Error searching images in PDF: {e}")
            return []

    @staticmethod
    def extract_from_docx(file_bytes: bytes) -> str:
        """
        Extract text, headings, and tables from Word (.docx) document.
        """
        try:
            doc = Document(io.BytesIO(file_bytes))
            paragraphs = []
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text)

            # Extract data from tables
            for t_idx, table in enumerate(doc.tables):
                table_lines = [f"\n[Table {t_idx+1}]"]
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    table_lines.append(" | ".join(row_cells))
                paragraphs.append("\n".join(table_lines))

            return "\n\n".join(paragraphs).strip()
        except Exception as e:
            logger.error(f"Error parsing DOCX: {e}")
            raise RuntimeError(f"Failed to read DOCX: {e}")

    @staticmethod
    def validate_image(file_bytes: bytes) -> Tuple[bool, Optional[str], Optional[Tuple[int, int]]]:
        """
        Validate image, returning (is_valid, format, dimensions).
        """
        try:
            with Image.open(io.BytesIO(file_bytes)) as img:
                return True, img.format, img.size
        except Exception as e:
            logger.warning(f"Error validating image: {e}")
            return False, None, None
