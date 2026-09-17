import io
import os
import pytest
import asyncio
from pathlib import Path
from docx import Document
from pypdf import PdfWriter
from PIL import Image
from bot.services.doc_parser import DocParser
from bot.services.storage_service import StorageService
from bot.services.system_status import SystemStatusService

@pytest.mark.asyncio
async def test_storage_service(tmp_path):
    # Configure temporary paths
    cloud_dir = tmp_path / "cloud"
    notes_dir = tmp_path / "notes"
    cloud_dir.mkdir(parents=True, exist_ok=True)
    notes_dir.mkdir(parents=True, exist_ok=True)
    service = StorageService(cloud_dir=cloud_dir, notes_dir=notes_dir)

    # 1. Test cloud storage saving
    dummy_data = b"Hello Cloud Storage!"
    res_cloud = await service.save_to_cloud(dummy_data, "test_file.txt")
    assert Path(res_cloud["path"]).exists()
    assert res_cloud["size_kb"] > 0
    assert "test_file.txt" in res_cloud["filename"]

    # 2. Test markdown note creation (.md)
    res_note = await service.save_note(
        title="Оплата счета за интернет",
        content="## Задачи\n- [ ] Оплатить до 15 числа\n- [ ] Сумма: 500 руб",
        note_type="voice",
        tags=["счета", "интернет"],
        raw_text="Надо оплатить интернет до пятнадцатого пятьсот рублей"
    )
    note_path = Path(res_note["path"])
    assert note_path.exists()
    assert note_path.suffix == ".md"
    
    # Verify .md file content
    content = note_path.read_text(encoding="utf-8")
    assert 'title: "Оплата счета за интернет"' in content
    assert 'tags: ["счета", "интернет"]' in content
    assert "- [ ] Оплатить до 15 числа" in content
    assert "Исходный текст / расшифровка" not in content

def test_docx_parser():
    # Create test .docx in memory
    doc = Document()
    doc.add_heading("Договор аренды", level=1)
    doc.add_paragraph("Арендатор обязуется внести оплату в размере 50 000 руб.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Услуга"
    table.rows[0].cells[1].text = "Стоимость"
    table.rows[1].cells[0].text = "Аренда"
    table.rows[1].cells[1].text = "50000"

    bio = io.BytesIO()
    doc.save(bio)
    file_bytes = bio.getvalue()

    extracted = DocParser.extract_from_docx(file_bytes)
    assert "Договор аренды" in extracted
    assert "50 000 руб." in extracted
    assert "Услуга | Стоимость" in extracted

def test_image_validation():
    img = Image.new("RGB", (100, 100), color="blue")
    bio = io.BytesIO()
    img.save(bio, format="JPEG")
    img_bytes = bio.getvalue()

    is_valid, fmt, size = DocParser.validate_image(img_bytes)
    assert is_valid is True
    assert fmt == "JPEG"
    assert size == (100, 100)

def test_pdf_parser():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    bio = io.BytesIO()
    writer.write(bio)
    file_bytes = bio.getvalue()

    text, pages = DocParser.extract_from_pdf(file_bytes)
    assert pages == 1
    assert isinstance(text, str)

def test_system_status():
    status_msg = SystemStatusService.format_status_message()
    assert "Статус бота и системы" in status_msg
    assert "AI Движок" in status_msg
    assert "Хранилище" in status_msg

def test_config_env_path(tmp_path):
    config_dir = tmp_path / "data" / "config"
    config_dir.mkdir(parents=True)
    env_file = config_dir / ".env"
    env_file.write_text("DEFAULT_ACTION=translate\nTARGET_LANGUAGE=Немецкий\n", encoding="utf-8")

    from config import Settings
    custom_settings = Settings(_env_file=str(env_file))
    assert custom_settings.DEFAULT_ACTION == "translate"
    assert custom_settings.TARGET_LANGUAGE == "Немецкий"

@pytest.mark.asyncio
async def test_web_server_endpoints():
    from bot.services.web_server import handle_dashboard, handle_health
    from aiohttp import web
    
    # Dashboard HTML
    req = web.Request
    resp_dashboard = await handle_dashboard(None)
    assert resp_dashboard.status == 200
    assert resp_dashboard.content_type == "text/html"
    assert "Telegram Helper Bot" in resp_dashboard.text

    # Health JSON
    resp_health = await handle_health(None)
    assert resp_health.status == 200
    assert resp_health.content_type == "application/json"

def test_telegram_helpers_split_text():
    from bot.utils.telegram_helpers import split_text
    # Small text
    assert split_text("Hello") == ["Hello"]

    # Long text with paragraphs
    paragraphs = ["Paragraph " + str(i) * 100 for i in range(50)]
    long_text = "\n\n".join(paragraphs)
    chunks = split_text(long_text, max_chunk_size=500)
    assert len(chunks) > 1
    for c in chunks:
        assert len(c) <= 500

@pytest.mark.asyncio
async def test_safe_edit_text_fallback():
    from unittest.mock import AsyncMock
    from aiogram.exceptions import TelegramBadRequest
    from bot.utils.telegram_helpers import safe_edit_text

    mock_msg = AsyncMock()
    # First call with Markdown raises "can't parse entities"
    # Second call with parse_mode=None succeeds
    mock_msg.edit_text.side_effect = [
        TelegramBadRequest(method=AsyncMock(), message="Bad Request: can't parse entities: unclosed entity"),
        AsyncMock(text="plain text response")
    ]

    res = await safe_edit_text(mock_msg, "Some text with _unclosed entity", parse_mode="Markdown")
    assert mock_msg.edit_text.call_count == 2
    # Verify second call had parse_mode=None
    assert mock_msg.edit_text.call_args_list[1].kwargs["parse_mode"] is None

def test_pdf_extract_images():
    from bot.services.doc_parser import DocParser
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    bio = io.BytesIO()
    writer.write(bio)
    file_bytes = bio.getvalue()

    images = DocParser.extract_images_from_pdf(file_bytes)
    assert isinstance(images, list)

def test_localization_texts_complete():
    from bot.texts import TEXTS, SUPPORTED_LANGUAGES, get_text, get_target_language_name, ALL_MENU_BUTTONS
    
    assert set(SUPPORTED_LANGUAGES.keys()) == {"ru", "en", "he"}
    
    # Check that keys are consistent across languages
    ru_keys = set(TEXTS["ru"].keys())
    en_keys = set(TEXTS["en"].keys())
    he_keys = set(TEXTS["he"].keys())

    assert ru_keys == en_keys, f"Missing in EN: {ru_keys - en_keys}"
    assert ru_keys == he_keys, f"Missing in HE: {ru_keys - he_keys}"

    # Check get_text helper and fallback
    assert get_text("btn_translate", "ru") == "🌐 Перевод"
    assert get_text("btn_translate", "en") == "🌐 Translation"
    assert get_text("btn_translate", "he") == "🌐 תרגום"
    assert get_text("btn_translate", "unknown_lang") == "🌐 Перевод"

    # Check formatting
    assert "Alice" in get_text("welcome", "ru", name="Alice")
    assert "Alice" in get_text("welcome", "en", name="Alice")
    assert "Alice" in get_text("welcome", "he", name="Alice")

    # Target language names
    assert get_target_language_name("ru") == "Русский"
    assert get_target_language_name("en") == "English"
    assert get_target_language_name("he") == "עברית"

    # Button sets
    assert "🌐 Перевод" in ALL_MENU_BUTTONS
    assert "🌐 Translation" in ALL_MENU_BUTTONS
    assert "🌐 תרגום" in ALL_MENU_BUTTONS

def test_user_settings_persistence(tmp_path):
    from bot.services.user_settings import UserSettingsService

    config_file = tmp_path / "user_settings.json"
    service = UserSettingsService(config_file=config_file)

    # Default is 'ru' without asking
    assert service.get_language(12345) == "ru"
    assert service.get_language(None) == "ru"

    # Change to english
    service.set_language(12345, "en")
    assert service.get_language(12345) == "en"

    # Change to hebrew
    service.set_language(67890, "he")
    assert service.get_language(67890) == "he"

    # Reload from disk
    service2 = UserSettingsService(config_file=config_file)
    assert service2.get_language(12345) == "en"
    assert service2.get_language(67890) == "he"
    assert service2.get_language(99999) == "ru"

    # Invalid language falls back to 'ru'
    service.set_language(12345, "invalid_lang")
    assert service.get_language(12345) == "ru"

def test_localized_keyboards():
    from bot.keyboards.reply import get_main_menu_keyboard
    from bot.keyboards.inline import get_cancel_keyboard, get_media_actions_keyboard, get_language_keyboard

    # Reply keyboard buttons
    kb_ru = get_main_menu_keyboard("ru")
    assert any("Анализ и перевод" in btn.text for row in kb_ru.keyboard for btn in row)
    assert any("Сканирование" in btn.text for row in kb_ru.keyboard for btn in row)
    assert any("Язык" in btn.text for row in kb_ru.keyboard for btn in row)

    kb_en = get_main_menu_keyboard("en")
    assert any("Analyze & Translate" in btn.text for row in kb_en.keyboard for btn in row)
    assert any("Scan text" in btn.text for row in kb_en.keyboard for btn in row)
    assert any("Language" in btn.text for row in kb_en.keyboard for btn in row)

    kb_he = get_main_menu_keyboard("he")
    assert any("ניתוח ותרגום" in btn.text for row in kb_he.keyboard for btn in row)
    assert any("סריקה" in btn.text for row in kb_he.keyboard for btn in row)
    assert any("שפה" in btn.text for row in kb_he.keyboard for btn in row)

    # Cancel keyboard
    assert "Отмена" in get_cancel_keyboard("ru").inline_keyboard[0][0].text
    assert "Cancel" in get_cancel_keyboard("en").inline_keyboard[0][0].text
    assert "ביטול" in get_cancel_keyboard("he").inline_keyboard[0][0].text

    # Media actions keyboard
    kb_act_analyze = get_media_actions_keyboard(file_type="image", current_action="analyze", lang="ru")
    datas_analyze = [b.callback_data for row in kb_act_analyze.inline_keyboard for b in row]
    assert "act_translate" in datas_analyze
    assert "act_scan" in datas_analyze
    assert "act_analyze" not in datas_analyze

    kb_act_scan = get_media_actions_keyboard(file_type="image", current_action="scan", lang="ru")
    datas_scan = [b.callback_data for row in kb_act_scan.inline_keyboard for b in row]
    assert "act_translate" in datas_scan
    assert "act_analyze" in datas_scan
    assert "act_scan" not in datas_scan

    # Language selection keyboard
    lang_kb = get_language_keyboard("en")
    buttons_flat = [b for row in lang_kb.inline_keyboard for b in row]
    callbacks = [b.callback_data for b in buttons_flat]
    assert "lang_set:ru" in callbacks
    assert "lang_set:en" in callbacks
    assert "lang_set:he" in callbacks
    # English should have checkmark
    en_btn = next(b for b in buttons_flat if b.callback_data == "lang_set:en")
    assert "✓" in en_btn.text

def test_system_status_multilingual():
    from bot.services.system_status import SystemStatusService

    msg_ru = SystemStatusService.format_status_message(lang="ru")
    assert "Статус бота и системы" in msg_ru
    assert "Русский" in msg_ru

    msg_en = SystemStatusService.format_status_message(lang="en")
    assert "Bot & System Status" in msg_en
    assert "English" in msg_en

    msg_he = SystemStatusService.format_status_message(lang="he")
    assert "סטטוס בוט ומערכת" in msg_he
    assert "עברית" in msg_he

@pytest.mark.asyncio
async def test_ai_service_prompts_multilingual(monkeypatch):
    from bot.services.ai_service import AIService
    
    captured_prompts = []
    async def fake_generate_text(prompt: str) -> str:
        captured_prompts.append(prompt)
        return '{"title": "Test Title", "tags": ["tag"], "content": "Content"}'

    service = AIService()
    monkeypatch.setattr(service, "_generate_text", fake_generate_text)

    # 1. Translate in Russian (default)
    await service.translate_text("Hello world", lang="ru")
    assert "Русский" in captured_prompts[-1]
    assert "Hello world" in captured_prompts[-1]

    # 2. Translate in English
    await service.translate_text("Привет мир", lang="en")
    assert "English" in captured_prompts[-1]
    assert "Привет мир" in captured_prompts[-1]

    # 3. Translate in Hebrew
    await service.translate_text("Hello world", lang="he")
    assert "עברית" in captured_prompts[-1]
    assert "Hello world" in captured_prompts[-1]

    # 4. Document Analysis in Russian
    await service.analyze_document_text("Счет на оплату", lang="ru")
    assert "Тип документа" in captured_prompts[-1]

    # 5. Document Analysis in English
    await service.analyze_document_text("Invoice 123", lang="en")
    assert "Document Type" in captured_prompts[-1]

    # 6. Document Analysis in Hebrew
    await service.analyze_document_text("חשבונית 123", lang="he")
    assert "סוג המסמך" in captured_prompts[-1]

def test_gemini_models_chain_config():
    from config import Settings
    
    # 1. Custom model and custom fallbacks
    s1 = Settings(
        GEMINI_MODEL="gemini-3.8-flash",
        GEMINI_FALLBACK_MODELS="gemini-3.7-flash,gemini-3.6-flash"
    )
    chain1 = s1.gemini_models_chain
    assert chain1[0] == "gemini-3.8-flash"
    assert "gemini-3.7-flash" in chain1
    assert "gemini-3.6-flash" in chain1
    assert len(chain1) == len(set(chain1))  # No duplicates

    # 2. Duplicate entries stripped
    s2 = Settings(
        GEMINI_MODEL="gemini-3.8-flash",
        GEMINI_FALLBACK_MODELS="gemini-3.8-flash,gemini-3.7-flash"
    )
    chain2 = s2.gemini_models_chain
    assert chain2[0] == "gemini-3.8-flash"
    assert chain2.count("gemini-3.8-flash") == 1

@pytest.mark.asyncio
async def test_gemini_fallback_on_503_error(monkeypatch):
    from unittest.mock import MagicMock
    from bot.services.ai_service import AIService
    from config import settings

    service = AIService()
    mock_client = MagicMock()
    service.gemini_client = mock_client
    service.provider = "gemini"

    calls = []
    def fake_generate_content(model, contents):
        calls.append(model)
        if model == settings.gemini_models_chain[0]:
            raise Exception("503 UNAVAILABLE: This model is currently experiencing high demand.")
        res = MagicMock()
        res.text = f"Success from {model}"
        return res

    mock_client.models.generate_content.side_effect = fake_generate_content

    result = await service._call_gemini_with_fallback("Test prompt")
    
    # Verify fallback was used
    assert len(calls) == 2
    assert calls[0] == settings.gemini_models_chain[0]
    assert calls[1] == settings.gemini_models_chain[1]
    assert result == f"Success from {settings.gemini_models_chain[1]}"

@pytest.mark.asyncio
async def test_gemini_fallback_all_fail(monkeypatch):
    from unittest.mock import MagicMock
    from bot.services.ai_service import AIService
    from config import settings

    service = AIService()
    mock_client = MagicMock()
    service.gemini_client = mock_client
    service.provider = "gemini"

    mock_client.models.generate_content.side_effect = Exception("Service error")

    with pytest.raises(Exception) as exc_info:
        await service._call_gemini_with_fallback("Test prompt")
    
    assert "Service error" in str(exc_info.value)
    # Check that all models in the chain were attempted
    assert mock_client.models.generate_content.call_count == len(settings.gemini_models_chain)

def test_allowed_users_configuration(monkeypatch):
    from config import Settings
    # 1. Test ALLOWED_USERS with names
    s1 = Settings(ALLOWED_USERS="1001:Иван, 1002:Мария", ALLOWED_USER_IDS="")
    assert s1.allowed_users == [1001, 1002]
    assert s1.allowed_users_map == {1001: "Иван", 1002: "Мария"}
    assert s1.get_user_name(1001) == "Иван"
    assert s1.get_user_name(1002) == "Мария"
    assert s1.get_user_name(9999) == "User 9999"

    # 2. Test fallback to ALLOWED_USER_IDS
    s2 = Settings(ALLOWED_USERS="", ALLOWED_USER_IDS="2001, 2002")
    assert s2.allowed_users == [2001, 2002]
    assert s2.get_user_name(2001) == "User 2001"

@pytest.mark.asyncio
async def test_per_user_storage_and_sharing(tmp_path):
    cloud_dir = tmp_path / "cloud"
    notes_dir = tmp_path / "notes"
    service = StorageService(cloud_dir=cloud_dir, notes_dir=notes_dir)

    # 1. Test saving file into user-specific folder
    file_data = b"User file content"
    res_cloud = await service.save_to_cloud(file_data, "doc.pdf", user_id=12345)
    cloud_file_path = Path(res_cloud["path"])
    assert cloud_file_path.exists()
    assert cloud_file_path.parent == cloud_dir / "12345"

    # 2. Test saving note into user-specific folder
    res_note = await service.save_note(
        title="План на неделю",
        content="Купить сервер",
        note_type="text",
        tags=["планы"],
        user_id=12345
    )
    note_path = Path(res_note["path"])
    assert note_path.exists()
    assert note_path.parent == notes_dir / "12345"
    content = note_path.read_text(encoding="utf-8")
    assert 'author_id: "12345"' in content
    token = res_note["token"]
    assert token is not None

    # Retrieve from token cache
    cached = service.get_note_by_token(token)
    assert cached["title"] == "План на неделю"

    # 3. Test sharing note
    share_result = await service.share_note(
        token_or_path=token,
        sender_id=12345,
        recipient_id=67890,
        sender_name="Алексей",
        recipient_name="Мария"
    )
    shared_path = Path(share_result["path"])
    assert shared_path.exists()
    assert shared_path.parent == notes_dir / "shared"
    assert "from_Алексей_" in shared_path.name

    shared_text = shared_path.read_text(encoding="utf-8")
    assert 'author: "Алексей"' in shared_text
    assert 'shared_to: "Мария"' in shared_text
    assert "Купить сервер" in shared_text

    # 4. Verify get_stats counts files across subdirectories and shared folder
    stats = service.get_stats()
    assert stats["cloud_files_count"] == 1
    assert stats["notes_count"] == 2  # 1 personal + 1 shared

def test_share_keyboards(monkeypatch):
    from config import settings
    from bot.keyboards.inline import get_note_share_keyboard, get_recipients_keyboard

    monkeypatch.setattr(settings, "ALLOWED_USERS", "111:Иван, 222:Мария")

    # Share button
    share_kb = get_note_share_keyboard("tok123", lang="ru")
    assert any(btn.callback_data == "share_start:tok123" for row in share_kb.inline_keyboard for btn in row)

    # Recipients keyboard for sender 111 (should only show 222: Мария)
    recipients_kb = get_recipients_keyboard("tok123", current_user_id=111, lang="ru")
    callback_datas = [btn.callback_data for row in recipients_kb.inline_keyboard for btn in row]
    button_texts = [btn.text for row in recipients_kb.inline_keyboard for btn in row]

    assert "share_send:222:tok123" in callback_datas
    assert any("Мария" in t for t in button_texts)
    assert not any("Иван" in t for t in button_texts)  # Sender himself excluded

def test_user_manager_crud_and_persistence(tmp_path):
    from bot.services.user_manager import UserManagerService
    users_file = tmp_path / "users.json"
    mgr = UserManagerService(config_file=users_file)

    # Initially empty
    assert mgr.get_all_users() == {}

    # Add admin user
    mgr.add_user(user_id=101, name="AdminAlex", role="admin", username="alex")
    assert mgr.is_allowed(101) is True
    assert mgr.is_admin(101) is True
    assert mgr.get_admin_ids() == [101]
    assert mgr.get_user_name(101) == "AdminAlex"

    # Add regular user
    mgr.add_user(user_id=102, name="Maria", role="user")
    assert mgr.is_allowed(102) is True
    assert mgr.is_admin(102) is False
    assert mgr.is_allowed(999) is False

    # Check mapping
    users_map = mgr.get_allowed_users_map()
    assert users_map == {101: "AdminAlex", 102: "Maria"}

    # Update name
    mgr.update_user_name(102, "Maria New")
    assert mgr.get_user_name(102) == "Maria New"

    # Verify reload from disk
    mgr_reloaded = UserManagerService(config_file=users_file)
    assert mgr_reloaded.get_user_name(102) == "Maria New"
    assert mgr_reloaded.is_admin(101) is True

    # Remove user
    mgr.remove_user(102)
    assert mgr.is_allowed(102) is False
    assert 102 not in mgr.get_allowed_user_ids()

def test_admin_keyboards():
    from bot.keyboards.inline import (
        get_request_access_keyboard,
        get_admin_request_keyboard,
        get_admin_main_keyboard,
        get_admin_users_list_keyboard,
        get_admin_user_card_keyboard
    )

    req_kb = get_request_access_keyboard(lang="ru")
    assert any(b.callback_data == "req_access" for row in req_kb.inline_keyboard for b in row)

    adm_req_kb = get_admin_request_keyboard(12345, "Иван", lang="ru")
    datas = [b.callback_data for row in adm_req_kb.inline_keyboard for b in row]
    assert "adm_appr:12345" in datas
    assert "adm_rejc:12345" in datas

    main_kb = get_admin_main_keyboard(lang="ru")
    main_datas = [b.callback_data for row in main_kb.inline_keyboard for b in row]
    assert "adm_list" in main_datas
    assert "adm_add" in main_datas

    users_dict = {
        "101": {"name": "Alex", "role": "admin"},
        "102": {"name": "Maria", "role": "user"}
    }
    list_kb = get_admin_users_list_keyboard(users_dict, lang="ru")
    list_datas = [b.callback_data for row in list_kb.inline_keyboard for b in row]
    assert "adm_view:101" in list_datas
    assert "adm_view:102" in list_datas

    card_kb = get_admin_user_card_keyboard(102, is_self=False, lang="ru")
    card_datas = [b.callback_data for row in card_kb.inline_keyboard for b in row]
    assert "adm_ren:102" in card_datas
    assert "adm_del:102" in card_datas

@pytest.mark.asyncio
async def test_auth_middleware(tmp_path, monkeypatch):
    from unittest.mock import AsyncMock, MagicMock
    from bot.middlewares.auth import AuthMiddleware
    from bot.services.user_manager import UserManagerService

    # Use isolated UserManager
    test_mgr = UserManagerService(config_file=tmp_path / "test_users.json")
    test_mgr.add_user(101, "AllowedUser", role="admin")
    monkeypatch.setattr("bot.middlewares.auth.user_manager", test_mgr)

    middleware = AuthMiddleware()
    next_handler = AsyncMock(return_value="OK")

    # 1. Allowed user message
    event_allowed = MagicMock()
    event_allowed.from_user.id = 101
    res = await middleware(next_handler, event_allowed, {})
    assert res == "OK"

    # 2. Unauthorized user message
    event_unauthorized = MagicMock()
    event_unauthorized.from_user.id = 999
    event_unauthorized.answer = AsyncMock()
    res_block = await middleware(next_handler, event_unauthorized, {})
    assert res_block is None
    assert event_unauthorized.answer.called

    # 3. Unauthorized user req_access callback query
    event_callback = MagicMock()
    event_callback.from_user.id = 999
    event_callback.data = "req_access"
    res_cb = await middleware(next_handler, event_callback, {})
    assert res_cb == "OK"


@pytest.mark.asyncio
async def test_auth_middleware_first_user_bootstrap(tmp_path, monkeypatch):
    from unittest.mock import AsyncMock, MagicMock
    from bot.middlewares.auth import AuthMiddleware
    from bot.services.user_manager import UserManagerService

    fresh_mgr = UserManagerService(config_file=tmp_path / "fresh_users.json")
    monkeypatch.setattr("bot.middlewares.auth.user_manager", fresh_mgr)
    monkeypatch.setattr("bot.handlers.admin.user_manager", fresh_mgr)

    middleware = AuthMiddleware()
    next_handler = AsyncMock(return_value="OK")

    # First user triggers bootstrap as admin
    event_first = MagicMock()
    event_first.from_user.id = 555
    event_first.from_user.first_name = "Creator"
    event_first.from_user.last_name = "Admin"
    event_first.from_user.username = "creator_bot"

    res = await middleware(next_handler, event_first, {})
    assert res == "OK"
    assert fresh_mgr.is_allowed(555) is True
    assert fresh_mgr.is_admin(555) is True
    assert fresh_mgr.get_admin_ids() == [555]

    # Second user is unauthorized and auto-generates access request
    event_second = MagicMock()
    event_second.from_user.id = 777
    event_second.from_user.first_name = "Newbie"
    event_second.from_user.last_name = ""
    event_second.from_user.username = "newbie77"
    event_second.answer = AsyncMock()

    mock_bot = AsyncMock()
    mock_bot.send_message = AsyncMock()

    res2 = await middleware(next_handler, event_second, {"bot": mock_bot})
    assert res2 is None
    assert event_second.answer.called
    assert mock_bot.send_message.called


@pytest.mark.asyncio
async def test_media_group_middleware():
    from unittest.mock import AsyncMock, MagicMock
    from bot.middlewares.album import MediaGroupMiddleware

    middleware = MediaGroupMiddleware(latency=0.05)
    next_handler = AsyncMock(return_value="PROCESSED_ALBUM")

    # Three messages in album
    msg1 = MagicMock()
    msg1.media_group_id = "mg_123"
    msg2 = MagicMock()
    msg2.media_group_id = "mg_123"
    msg3 = MagicMock()
    msg3.media_group_id = "mg_123"

    # Simulate sequential arrivals
    task1 = asyncio.create_task(middleware(next_handler, msg1, {}))
    await asyncio.sleep(0.01)
    task2 = asyncio.create_task(middleware(next_handler, msg2, {}))
    await asyncio.sleep(0.01)
    task3 = asyncio.create_task(middleware(next_handler, msg3, {}))

    res1 = await task1
    res2 = await task2
    res3 = await task3

    assert res1 is None
    assert res2 is None
    assert res3 == "PROCESSED_ALBUM"
    assert next_handler.call_count == 1
    call_args = next_handler.call_args[0]
    data = call_args[1]
    assert "album" in data
    assert len(data["album"]) == 3


@pytest.mark.asyncio
async def test_ai_service_analyze_document_images_batch(monkeypatch):
    from unittest.mock import AsyncMock
    from bot.services.ai_service import AIService

    service = AIService()
    mock_generate = AsyncMock(return_value="Merged and deduplicated analysis result")
    monkeypatch.setattr(service, "_call_gemini_with_fallback", mock_generate)

    res = await service.analyze_document_images_batch(
        images_bytes=[b"img1", b"img2"],
        mime_types=["image/jpeg", "image/jpeg"],
        action="analyze",
        lang="ru"
    )

    assert res == "Merged and deduplicated analysis result"
    assert mock_generate.called
    parts = mock_generate.call_args[0][0]
    assert len(parts) == 3  # 2 image parts + 1 prompt string


@pytest.mark.asyncio
async def test_ai_service_structure_note_formats(monkeypatch):
    from unittest.mock import AsyncMock
    from bot.services.ai_service import AIService

    service = AIService()

    # 1. Test markdown checklists for shopping list
    mock_checklist_response = '```json\n{"title": "Список покупок", "tags": ["покупки"], "content": "- [ ] Картошка\\n- [ ] Сыр\\n- [ ] Молоко"}\n```'
    monkeypatch.setattr(service, "_generate_text", AsyncMock(return_value=mock_checklist_response))

    note = await service.structure_note("Подготовь список покупок: картошка, сыр, молоко")
    assert note["title"] == "Список покупок"
    assert "- [ ] Картошка" in note["content"]
    assert "Подготовь список покупок" not in note["content"]

    # 2. Test markdown table formatting
    mock_table_response = '{"title": "Таблица сравнения", "tags": ["работа"], "content": "| Пункт | Описание |\\n|---|---|\\n| 1 | Картошка |"}'
    monkeypatch.setattr(service, "_generate_text", AsyncMock(return_value=mock_table_response))

    note_table = await service.structure_note("Сделай мне таблицу: пункт 1 картошка")
    assert note_table["title"] == "Таблица сравнения"
    assert "| Пункт | Описание |" in note_table["content"]


@pytest.mark.asyncio
async def test_scan_feature_and_prompts(monkeypatch):
    from unittest.mock import AsyncMock
    from bot.services.ai_service import AIService
    from bot.texts import BUTTON_SCAN_ALL

    assert "📷 Сканирование" in BUTTON_SCAN_ALL
    assert "📷 Scan text" in BUTTON_SCAN_ALL
    assert "📷 סריקה" in BUTTON_SCAN_ALL

    service = AIService()
    captured_vision = []
    async def fake_generate_vision(prompt: str, image_bytes: bytes, mime_type: str) -> str:
        captured_vision.append((prompt, image_bytes, mime_type))
        return "Распознанный текст накладной"

    monkeypatch.setattr(service, "_generate_vision", fake_generate_vision)

    # 1. Scan single image
    res = await service.scan_image(b"fake_image_bytes", "image/jpeg", lang="ru")
    assert res == "Распознанный текст накладной"
    assert "оптическое распознавание текста (OCR)" in captured_vision[-1][0]
    assert "Не переводите" in captured_vision[-1][0]

    # 2. Batch album scan
    mock_batch = AsyncMock(return_value="Сквозной объединённый текст альбома")
    monkeypatch.setattr(service, "_call_gemini_with_fallback", mock_batch)
    res_batch = await service.analyze_document_images_batch(
        images_bytes=[b"img1", b"img2"],
        action="scan",
        lang="ru"
    )
    assert res_batch == "Сквозной объединённый текст альбома"
    prompt_used = mock_batch.call_args[0][0][-1]
    assert "серия из 2 изображений" in prompt_used
    assert "Не переводите" in prompt_used

    # 3. Multimodal scan with digital text
    res_pdf_text = await service.scan_document_multimodal(
        file_bytes=b"",
        text_content="Digital PDF text content",
        lang="ru"
    )
    assert res_pdf_text == "Digital PDF text content"


@pytest.mark.asyncio
async def test_ai_service_structure_note_leakage_fallback(monkeypatch):
    from unittest.mock import AsyncMock
    from bot.services.ai_service import AIService

    service = AIService()

    # 1. Simulate model hallucinating English prompt instructions
    mock_leaked_response_en = '''```json
{
  "title": "Voice Note Pending",
  "tags": ["task", "audio"],
  "content": "- [ ] Review pending voice transcription\\n- [ ] Update note with finalized transcription details"
}
```'''
    monkeypatch.setattr(service, "_generate_text", AsyncMock(return_value=mock_leaked_response_en))

    user_voice = "Завтра в 15:00 встреча с архитектором по проекту дома"
    note_en = await service.structure_note(user_voice, lang="en")
    
    # Must fallback to user's real voice text and not the prompt text
    assert note_en["content"] == user_voice
    assert "Review pending voice transcription" not in note_en["content"]
    assert note_en["title"] != "Voice Note Pending"

    # 2. Simulate model hallucinating Russian prompt instructions
    mock_leaked_response_ru = '''{
  "title": "Инструкция по обработке голосовых заметок",
  "tags": ["заметки", "структура"],
  "content": "- [ ] Выделить полезное содержимое и основные мысли\\n- [ ] Оформить задачи в виде чекбоксов\\n- [ ] Удалить вводные фразы"
}'''
    monkeypatch.setattr(service, "_generate_text", AsyncMock(return_value=mock_leaked_response_ru))

    user_text = "Купить корм для кота и забрать посылку"
    note_ru = await service.structure_note(user_text, lang="ru")

    assert note_ru["content"] == user_text
    assert "Выделить полезное содержимое" not in note_ru["content"]
    assert note_ru["title"] != "Инструкция по обработке голосовых заметок"


@pytest.mark.asyncio
async def test_ai_service_extract_doc_note_meta(monkeypatch):
    from unittest.mock import AsyncMock
    from bot.services.ai_service import AIService

    service = AIService()
    mock_meta_response = '{"title": "Договор аренды жилья", "tags": ["документ", "аренда"]}'
    monkeypatch.setattr(service, "_generate_text", AsyncMock(return_value=mock_meta_response))

    doc_text = "📋 **Тип документа**: Договор аренды жилого помещения\n🎯 **Краткая суть**: Аренда квартиры на год."
    meta = await service.extract_doc_note_meta(doc_text, lang="ru")

    assert meta["title"] == "Договор аренды жилья"
    assert "документ" in meta["tags"]
    assert "аренда" in meta["tags"]


@pytest.mark.asyncio
async def test_callback_save_note_preserves_full_content(tmp_path, monkeypatch):
    from unittest.mock import AsyncMock, MagicMock
    from bot.handlers.actions import callback_save_note
    from bot.services.storage_service import StorageService

    # Use isolated StorageService
    test_storage = StorageService(cloud_dir=tmp_path / "cloud", notes_dir=tmp_path / "notes")
    monkeypatch.setattr("bot.handlers.actions.storage_service", test_storage)

    # Mock extract_doc_note_meta
    mock_extract_meta = AsyncMock(return_value={"title": "Анализ счёта за электричество", "tags": ["счета"]})
    monkeypatch.setattr("bot.handlers.actions.ai_service.extract_doc_note_meta", mock_extract_meta)

    # State has large multi-paragraph text (e.g., from translation or document analysis)
    full_document_text = (
        "📋 **Тип документа**: Квитанция на оплату электроэнергии\n"
        "🎯 **Краткая суть**: Сумма к оплате 3450 руб до 25 сентября.\n"
        "⚠️ **Что требуется от вас**: Оплатить по QR коду или в личном кабинете.\n"
        "🌐 **Перевод ключевых положений**: Начислено за август: 450 кВтч."
    )
    mock_state = AsyncMock()
    mock_state.get_data.return_value = {"last_extracted_text": full_document_text}

    mock_query = AsyncMock()
    mock_query.from_user.id = 12345
    mock_query.message.text = "Last chunk of text"
    mock_query.message.reply = AsyncMock()

    await callback_save_note(mock_query, mock_state, lang="ru")

    # Check that note was saved into storage
    notes = list((tmp_path / "notes" / "12345").glob("*.md"))
    assert len(notes) == 1
    saved_md = notes[0].read_text(encoding="utf-8")

    # Verify full content was preserved without loss
    assert full_document_text in saved_md
    assert 'title: "Анализ счёта за электричество"' in saved_md
    assert 'tags: ["документ", "счета"]' in saved_md








